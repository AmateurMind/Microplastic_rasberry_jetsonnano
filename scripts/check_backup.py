import docx

doc = docx.Document("Final BE Project Stage I Report_BACKUP.docx")
print("Total paragraphs:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        print(f"[{i:3d}] {txt}")
